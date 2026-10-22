"""HTTP-to-worker document proof for an isolated five-service staging stack."""

import io
import os
import time
from urllib.parse import urlparse

import httpx
import pytest
from pptx import Presentation

from app.database import SessionLocal
from app.models.content import UploadedDocument
from app.models.embedding import Embedding
from app.models.subject import SubjectPackage, UserSubjectAccess
from app.models.user import User


API_URL = os.getenv("STAGING_API_URL")
pytestmark = pytest.mark.skipif(not API_URL, reason="STAGING_API_URL is not configured")


def _assert_isolated_runtime() -> None:
    parsed = urlparse(API_URL)
    if parsed.hostname not in {"127.0.0.1", "localhost", "backend"}:
        pytest.fail("STAGING_API_URL must target localhost or the isolated backend service")
    db_host = urlparse(os.environ.get("DATABASE_URL", "")).hostname
    if db_host not in {"127.0.0.1", "localhost", "db"}:
        pytest.fail("DATABASE_URL must target localhost or the isolated db service")


def _minimal_pdf(text: str) -> bytes:
    safe = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
    stream = f"BT /F1 12 Tf 72 720 Td ({safe}) Tj ET".encode("ascii")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 5 0 R >> >> /Contents 4 0 R >>",
        b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    payload = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for number, obj in enumerate(objects, 1):
        offsets.append(len(payload))
        payload.extend(f"{number} 0 obj\n".encode("ascii") + obj + b"\nendobj\n")
    xref = len(payload)
    payload.extend(f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode("ascii"))
    for offset in offsets[1:]:
        payload.extend(f"{offset:010d} 00000 n \n".encode("ascii"))
    payload.extend(
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode("ascii")
    )
    return bytes(payload)


def _pptx_bytes(text: str) -> bytes:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "Nuvira synthetic staging material"
    slide.placeholders[1].text = text
    output = io.BytesIO()
    presentation.save(output)
    return output.getvalue()


def _register(client: httpx.Client, email: str) -> str:
    response = client.post(
        "/auth/register",
        json={"email": email, "password": "SyntheticPass1234", "full_name": "Synthetic Student"},
    )
    assert response.status_code == 201, response.text
    return response.json()["access_token"]


def _wait_for_document(client: httpx.Client, token: str, subject_id: int, document_id: int) -> dict:
    deadline = time.monotonic() + 100
    while time.monotonic() < deadline:
        response = client.get(
            "/agent/materials",
            params={"subject_package_id": subject_id},
            headers={"Authorization": f"Bearer {token}"},
        )
        assert response.status_code == 200, response.text
        document = next(item for item in response.json()["documents"] if item["id"] == document_id)
        if document["status"] == "ready":
            return document
        time.sleep(2)
    pytest.fail(f"document {document_id} did not reach ready status")


def test_synthetic_pdf_pptx_worker_processing_and_student_isolation():
    _assert_isolated_runtime()
    unique = str(time.time_ns())
    email_a = f"staging-a-{unique}@example.com"
    email_b = f"staging-b-{unique}@example.com"
    with httpx.Client(base_url=API_URL, timeout=20) as client_a, httpx.Client(base_url=API_URL, timeout=20) as client_b:
        token_a = _register(client_a, email_a)
        token_b = _register(client_b, email_b)

        db = SessionLocal()
        try:
            user_a = db.query(User).filter(User.email == email_a).one()
            user_b = db.query(User).filter(User.email == email_b).one()
            subject = SubjectPackage(name=f"Synthetic Chemistry {unique}", language="en")
            db.add(subject)
            db.flush()
            db.add_all([
                UserSubjectAccess(user_id=user_a.id, subject_package_id=subject.id, source="gifted"),
                UserSubjectAccess(user_id=user_b.id, subject_package_id=subject.id, source="gifted"),
            ])
            db.commit()
            subject_id = subject.id
            user_a_id = user_a.id
        finally:
            db.close()

        headers_a = {"Authorization": f"Bearer {token_a}"}
        pdf_response = client_a.post(
            "/agent/materials/upload",
            data={"subject_package_id": str(subject_id)},
            files={"file": ("synthetic.pdf", _minimal_pdf("Synthetic chemistry PDF content"), "application/pdf")},
            headers=headers_a,
        )
        assert pdf_response.status_code == 200, pdf_response.text
        pdf_id = pdf_response.json()["document_id"]
        assert _wait_for_document(client_a, token_a, subject_id, pdf_id)["chunk_count"] > 0

        pptx_response = client_a.post(
            "/agent/materials/upload",
            data={"subject_package_id": str(subject_id)},
            files={"file": (
                "transient.pptx",
                _pptx_bytes("NUVIRA_TRANSIENT_EMBEDDING_FAILURE then recover safely"),
                "application/vnd.openxmlformats-officedocument.presentationml.presentation",
            )},
            headers=headers_a,
        )
        assert pptx_response.status_code == 200, pptx_response.text
        pptx_id = pptx_response.json()["document_id"]
        assert _wait_for_document(client_a, token_a, subject_id, pptx_id)["chunk_count"] > 0

        b_materials = client_b.get(
            "/agent/materials",
            params={"subject_package_id": subject_id},
            headers={"Authorization": f"Bearer {token_b}"},
        )
        assert b_materials.status_code == 200
        assert b_materials.json()["documents"] == []

        unauthenticated = client_a.get("/agent/materials", params={"subject_package_id": subject_id})
        assert unauthenticated.status_code == 401

        malformed = client_a.post(
            "/agent/materials/upload",
            data={"subject_package_id": str(subject_id)},
            files={"file": ("malformed.pdf", b"not-a-pdf", "application/pdf")},
            headers=headers_a,
        )
        assert malformed.status_code == 400
        unsupported = client_a.post(
            "/agent/materials/upload",
            data={"subject_package_id": str(subject_id)},
            files={"file": ("notes.txt", b"synthetic", "text/plain")},
            headers=headers_a,
        )
        assert unsupported.status_code == 400
        oversized = client_a.post(
            "/agent/materials/upload",
            data={"subject_package_id": str(subject_id)},
            files={"file": ("oversized.pdf", b"%PDF-" + b"x" * 1_048_576, "application/pdf")},
            headers=headers_a,
        )
        assert oversized.status_code == 413

    db = SessionLocal()
    try:
        documents = db.query(UploadedDocument).filter(UploadedDocument.id.in_([pdf_id, pptx_id])).all()
        assert len(documents) == 2
        assert all(document.uploaded_by_user_id == user_a_id for document in documents)
        assert next(document for document in documents if document.id == pptx_id).processing_attempts == 2
        embeddings = db.query(Embedding).filter(Embedding.document_id.in_([pdf_id, pptx_id])).all()
        assert embeddings
        assert all(row.user_id == user_a_id for row in embeddings)
    finally:
        db.close()
