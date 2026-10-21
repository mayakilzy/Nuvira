"""
Document processing — استخراج نص من PDF/PPTX، تقطيع لقطع صغيرة، تحويل كل قطعة
لـ embedding مرتبط بالفصل والمادة، مع تحديث حالة الملف بقاعدة البيانات.
"""
from sqlalchemy.orm import Session

from app.models.content import UploadedDocument
from app.services import embedding_service, vector_service


def extract_text(path: str) -> str:
    if path.lower().endswith(".pdf"):
        from pypdf import PdfReader
        reader = PdfReader(path)
        return "\n".join(page.extract_text() or "" for page in reader.pages)
    elif path.lower().endswith((".pptx", ".ppt")):
        from pptx import Presentation
        prs = Presentation(path)
        chunks = []
        for slide in prs.slides:
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text:
                    chunks.append(shape.text)
        return "\n".join(chunks)
    else:
        raise ValueError(f"Unsupported file type: {path}")


def chunk_text(text: str, chunk_size: int = 800, overlap: int = 100) -> list[str]:
    """تقطيع بالكلمات مع تداخل بسيط، حتى لا يُقطَع مفهوم بمنتصف جملة مهمة."""
    words = text.split()
    chunks = []
    start = 0
    while start < len(words):
        end = start + chunk_size
        chunks.append(" ".join(words[start:end]))
        start = end - overlap
    return [c.strip() for c in chunks if c.strip()]


async def process_document(db: Session, document_id: int):
    doc = db.query(UploadedDocument).filter(UploadedDocument.id == document_id).first()
    if not doc:
        return

    try:
        text = extract_text(doc.storage_path)
        chunks = chunk_text(text)

        for chunk in chunks:
            vector = await embedding_service.get_embedding(chunk)
            # الربط بالسياق يتم في نفس الـ INSERT — بلا إعادة استعلام هشّة لأخر صف
            vector_service.add_embedding(
                db, content=chunk, vector=vector,
                user_id=doc.uploaded_by_user_id,
                meta={"filename": doc.filename},
                source_ref=doc.filename,
                subject_package_id=doc.subject_package_id,
                chapter_id=doc.chapter_id,
                document_id=doc.id,
            )

        doc.status = "ready"
        doc.chunk_count = len(chunks)
        db.commit()

    except Exception as e:
        doc.status = "failed"
        doc.error_note = str(e)[:500]
        db.commit()
